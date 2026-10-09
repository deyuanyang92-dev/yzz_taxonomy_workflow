"""Excel 报表统一格式: 粗体表头、冻结首行、筛选、列宽; 可按规则给整行着色."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
import pandas as pd

GREY = PatternFill("solid", fgColor="EEEEEE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def write_workbook(path, sheets, freeze=None, shade=None):
    """sheets: {name: DataFrame}; freeze: {sheet: 'C2'}; shade: {sheet: [(column, value_or_fn, fill)]}.
    目标文件被占用 (如在 Excel 中打开) 时另存为 <name>_HHMM.xlsx, 返回实际路径."""
    import datetime
    from pathlib import Path
    path = Path(path)
    try:
        open(path, "ab").close()
    except PermissionError:
        path = path.with_name(f"{path.stem}_{datetime.datetime.now():%H%M}{path.suffix}")
        print(f"   (原文件被占用, 另存为 {path.name})")
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in sheets.items():
            if df is not None and len(df.columns):
                df.to_excel(w, sheet_name=name[:31], index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():          # data cells stay text: no '=...' formulas from GenBank/user values
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"
        ws.freeze_panes = (freeze or {}).get(ws.title, "A2")
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
            ws.column_dimensions[col[0].column_letter].width = min(max(8, width + 2), 50)
        rules = (shade or {}).get(ws.title, [])
        if rules:
            hdr = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2):
                for col, test, fill in rules:
                    if col in hdr:
                        v = row[hdr.index(col)].value
                        if (test(v) if callable(test) else v == test):
                            for c in row:
                                c.fill = fill
    wb.save(path)
    return path
